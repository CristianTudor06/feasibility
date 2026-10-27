"""Convert an Excel capability database (.xlsx) into the CSV text the app's parser expects.

The UI parses databases with ``parseCSVToPitches`` (frontend/src/lib/csvParser.ts),
which splits on line breaks before honouring quotes and treats every ``"`` as a
quote toggle. Excel cells often contain both (long audience descriptions), so the
text is flattened here: line breaks become spaces and double quotes become single
quotes. Everything else is passed through for the existing parser to interpret.
"""

import csv
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

HEADER_HINTS = ("market", "country")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = re.sub(r"[\r\n]+", " ", str(value)).replace('"', "'")
    return re.sub(r"\s{2,}", " ", text).strip()


def _pick_sheet(wb):
    """The first sheet whose header row names a market/country column, else the first sheet."""
    for ws in wb.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if any(any(h in str(c or "").lower() for h in HEADER_HINTS) for c in header):
            return ws
    return wb.worksheets[0]


def xlsx_to_csv(stream):
    """Return (csv_text, sheet_name, data_row_count) for the capability sheet of a workbook."""
    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        ws = _pick_sheet(wb)
        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\n")
        count = 0
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if not cells:
                continue
            writer.writerow(cells)
            count += 1
        return out.getvalue(), ws.title, max(count - 1, 0)
    finally:
        wb.close()
