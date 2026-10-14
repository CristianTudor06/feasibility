"""Excel (.xlsx) export of the feasibility overview.

Replaces the original Google Sheets export: same columns and header/allocation
highlighting, but produced locally so it works on an internal network.
"""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "Pitch name", "Source", "Option", "Market", "Target", "LOI", "IR%", "IR universe", "Quotas",
    "Nº of Waves", "Total fresh sample", "Total Sample Size (CALCULATED)", "Internal Max", "Max Internal %",
    "Vendor 1 Cap", "Vendor 1 CPI", "Vendor 2 Cap", "Vendor 2 CPI", "Vendor 3 Cap", "Vendor 3 CPI",
    "Vendor 4 Cap", "Vendor 4 CPI", "Feasibility Status", "Allocated Internal", "Allocated Vendor 1",
    "Allocated Vendor 2", "Allocated Vendor 3", "Allocated Vendor 4", "Total Allocated Sample",
    "Unfilled Sample", "Max Potential", "Overlap Penalty %", "Internal Share %", "External Share %",
    "Used Vendors", "Weighted CPI ($)", "Number of Vendors", "PM Fee ($)", "Comparability Status",
    "Total Cost ($)", "OP - Operational Profit (%)",
]

# 0-based column index where the calculated allocation columns start ("Feasibility Status").
ALLOCATION_START = 22

NUMBER_FORMATS = {
    "IR%": "0%",
    "Max Internal %": "0%",
    "Overlap Penalty %": "0%",
    "Internal Share %": "0%",
    "External Share %": "0%",
    "Weighted CPI ($)": "0.00",
    "OP - Operational Profit (%)": "0.0%",
}

HEADER_FILL = PatternFill("solid", fgColor="1A3359")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
ALLOCATION_FILL = PatternFill("solid", fgColor="F0F7FF")


def _n(value, default=0):
    if value is None or isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        return value
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _row_values(row, alloc):
    sales_cpi = _n(row.get("salesCpi"), 7.5)
    revenue = _n(alloc.get("totalAllocated")) * sales_cpi
    op_costs = _n(alloc.get("externalCost")) + _n(alloc.get("internalCostRef"))
    op_ratio = (revenue - op_costs) / revenue if revenue > 0 else 0
    max_internal = row.get("maxInternalPercent")

    return [
        row.get("pitchName"),
        row.get("source"),
        row.get("option"),
        row.get("market"),
        row.get("target"),
        _n(row.get("loi")),
        _n(row.get("irPercent")),
        row.get("irUniverse"),
        row.get("quotas"),
        _n(row.get("waves")),
        _n(row.get("totalFreshSample")),
        _n(row.get("totalSampleSize")),
        _n(row.get("internal")),
        _n(max_internal, 100) / 100,
        _n(row.get("vendor1")),
        _n(row.get("vendor1Cpi")),
        _n(row.get("vendor2")),
        _n(row.get("vendor2Cpi")),
        _n(row.get("vendor3")),
        _n(row.get("vendor3Cpi")),
        _n(row.get("vendor4")),
        _n(row.get("vendor4Cpi")),
        alloc.get("status"),
        _n(alloc.get("allocatedInternal")),
        _n(alloc.get("allocatedVendor1")),
        _n(alloc.get("allocatedVendor2")),
        _n(alloc.get("allocatedVendor3")),
        _n(alloc.get("allocatedVendor4")),
        _n(alloc.get("totalAllocated")),
        _n(alloc.get("unfilledSample")),
        _n(alloc.get("maxPotential")),
        _n(alloc.get("overlapPenaltyPct")) / 100,
        _n(alloc.get("internalSharePct")) / 100,
        _n(alloc.get("externalSharePct")) / 100,
        alloc.get("activeVendorsCsv") or "None",
        round(_n(alloc.get("weightedCpi")), 2),
        _n(alloc.get("numVendorsUsed")),
        _n(alloc.get("pmFee")),
        alloc.get("comparability") or "High",
        _n(alloc.get("totalCost")),
        op_ratio,
    ]


def build_feasibility_workbook(pitches, allocations, strategy_name):
    alloc_by_row = {a.get("rowId"): a for a in allocations or [] if isinstance(a, dict)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Feasibility"
    wb.properties.title = (
        f"Pitch Feasibility Output ({(strategy_name or '').upper()}) - {datetime.now():%Y-%m-%d %H:%M}"
    )

    ws.append(HEADERS)
    for row in pitches or []:
        alloc = alloc_by_row.get(row.get("id")) or {
            "status": "Unfeasible",
            "unfilledSample": row.get("totalSampleSize"),
            "activeVendorsCsv": "None",
            "pmFee": 150,
            "comparability": "High",
        }
        ws.append(_row_values(row, alloc))

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        fmt = NUMBER_FORMATS.get(header)
        for r in range(2, ws.max_row + 1):
            data_cell = ws.cell(row=r, column=col_idx)
            if fmt:
                data_cell.number_format = fmt
            if col_idx - 1 >= ALLOCATION_START:
                data_cell.fill = ALLOCATION_FILL

        widest = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)) \
            if ws.max_row > 1 else 0
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(12, widest + 2, len(header) * 0.6), 40)

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
