import io

import pytest
from openpyxl import load_workbook

from feasibility import create_app

ROW = {
    "id": "row-1", "pitchName": "Pitch 1", "source": "Test", "option": "1", "market": "Germany",
    "target": "gen pop", "loi": 10, "irPercent": 0.3, "irUniverse": "gen pop", "quotas": "None",
    "waves": 1, "totalFreshSample": 1000, "totalSampleSize": 1000, "internal": 400,
    "vendor1": 500, "vendor1Cpi": 3, "vendor2": 500, "vendor2Cpi": 5,
    "vendor3": 0, "vendor3Cpi": 0, "vendor4": 0, "vendor4Cpi": 0, "salesCpi": 7.5,
}


@pytest.fixture
def client(tmp_path):
    app = create_app({"TESTING": True, "DATABASE": str(tmp_path / "test.db")})
    return app.test_client()


def test_capability_database_roundtrip(client):
    assert client.get("/api/database").get_json() is None

    res = client.put("/api/database", json={"rows": [ROW, {**ROW, "id": "row-2", "market": "France"}],
                                            "filename": "caps.csv"})
    assert res.status_code == 200

    db = client.get("/api/database").get_json()
    assert [r["market"] for r in db["rows"]] == ["Germany", "France"]
    assert db["filename"] == "caps.csv"
    assert db["uploadedAt"]

    # A new upload replaces the previous one entirely.
    client.put("/api/database", json={"rows": [ROW], "filename": "v2.csv"})
    assert len(client.get("/api/database").get_json()["rows"]) == 1

    assert client.delete("/api/database").status_code == 204
    assert client.get("/api/database").get_json() is None


def test_capability_database_rejects_bad_payload(client):
    res = client.put("/api/database", json={"rows": "nope"})
    assert res.status_code == 400
    assert "rows" in res.get_json()["error"]


def test_saved_options_crud(client):
    opt = {"id": "saved-1", "name": "Scenario A", "timestamp": "10:00 1/1/2026", "data": [ROW]}
    assert client.post("/api/saved-options", json=opt).status_code == 201
    assert client.post("/api/saved-options", json={**opt, "id": "saved-2", "name": "Scenario B"}).status_code == 201

    listed = client.get("/api/saved-options").get_json()
    assert [o["name"] for o in listed] == ["Scenario A", "Scenario B"]
    assert listed[0]["data"][0]["market"] == "Germany"

    # Names are unique (case-insensitive), as the UI expects.
    dup = client.post("/api/saved-options", json={**opt, "id": "saved-3", "name": "scenario a"})
    assert dup.status_code == 409

    updated = {**opt, "timestamp": "11:00 1/1/2026", "data": [{**ROW, "market": "Spain"}]}
    assert client.put("/api/saved-options/saved-1", json=updated).status_code == 200
    listed = client.get("/api/saved-options").get_json()
    assert listed[0]["data"][0]["market"] == "Spain"
    assert [o["id"] for o in listed] == ["saved-1", "saved-2"]  # order kept after an update

    assert client.put("/api/saved-options/missing", json=opt).status_code in (404, 409)
    assert client.delete("/api/saved-options/saved-1").status_code == 204
    assert [o["id"] for o in client.get("/api/saved-options").get_json()] == ["saved-2"]


def test_chat_requires_messages(client):
    res = client.post("/api/chat", json={})
    assert res.status_code == 400


def test_chat_cheapest_version_sets_strategy(client):
    res = client.post("/api/chat", json={
        "messages": [{"role": "user", "text": "create the cheapest version"}],
        "dataState": {"pitches": [{**ROW, "maxInternalPercent": 40}]},
        "activeStrategy": "tier",
    })
    body = res.get_json()
    assert res.status_code == 200
    assert body["updateStrategy"] == "cost"
    assert body["updatedPitches"][0]["maxInternalPercent"] == 100


@pytest.mark.parametrize("question, heading", [
    ("what is the loi rule?", "LOI (Duration) Calculation Rule"),
    ("explain the ir rule", "Incidence (IR%) Calculation Rule"),
    ("fara cote", "Quotas Rule"),
])
def test_chat_explains_matching_rules(client, question, heading):
    res = client.post("/api/chat", json={
        "messages": [{"role": "user", "text": question}],
        "dataState": {"pitches": [ROW]},
    })
    body = res.get_json()
    assert heading in body["text"]
    assert body["updatedPitches"] is None


def test_xlsx_export(client):
    alloc = {"rowId": "row-1", "status": "Fully Feasible", "allocatedInternal": 400, "allocatedVendor1": 500,
             "allocatedVendor2": 100, "totalAllocated": 1000, "unfilledSample": 0, "externalCost": 2000,
             "internalCostRef": 800, "totalCost": 2000, "weightedCpi": 3.3333, "internalSharePct": 40,
             "externalSharePct": 60, "numVendorsUsed": 2, "pmFee": 150, "activeVendorsCsv": "Vendor 1, Vendor 2"}
    res = client.post("/api/export/xlsx", json={"pitches": [ROW], "allocations": [alloc], "strategyName": "cost"})
    assert res.status_code == 200
    assert res.headers["Content-Disposition"].startswith("attachment")

    ws = load_workbook(io.BytesIO(res.data)).active
    headers = [c.value for c in ws[1]]
    assert headers[0] == "Pitch name" and headers[-1] == "OP - Operational Profit (%)"
    row = {h: c.value for h, c in zip(headers, ws[2])}
    assert row["Market"] == "Germany"
    assert row["Feasibility Status"] == "Fully Feasible"
    assert row["Weighted CPI ($)"] == 3.33
    # OP % = (revenue - (external + internal cost)) / revenue = (7500 - 2800) / 7500
    assert row["OP - Operational Profit (%)"] == pytest.approx((7500 - 2800) / 7500)


def test_unknown_api_route_returns_json(client):
    res = client.get("/api/does-not-exist")
    assert res.status_code == 404
    assert "error" in res.get_json()


def test_serves_frontend(client):
    res = client.get("/")
    assert res.status_code in (200, 503)  # 503 only when the frontend has not been built
